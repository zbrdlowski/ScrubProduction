from __future__ import annotations

import argparse
import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


DAY_CODES = {
    1: "MO",
    2: "TU",
    3: "WE",
    4: "TH",
    5: "FR",
    6: "SA",
    7: "SU",
}


def as_int(value: Any) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        return int(round(value))

    text = str(value).strip().replace("\xa0", "").replace(",", ".")
    if text == "":
        return None

    try:
        return int(round(float(text)))
    except ValueError:
        return None


def parse_stat_date(value: Any, year: int) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        return from_excel(value).date()

    text = str(value).strip()
    match = re.match(r"^(\d{1,2})\.(\d{1,2})\.?(?:(\d{2,4}))?$", text)
    if not match:
        return None

    day = int(match.group(1))
    month = int(match.group(2))
    parsed_year = int(match.group(3)) if match.group(3) else year
    if parsed_year < 100:
        parsed_year += 2000

    return date(parsed_year, month, day)


def extract_rows(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    rows: list[dict[str, Any]] = []

    for sheet_name in workbook.sheetnames:
        if not re.match(r"^\d{4}$", sheet_name):
            continue

        sheet_year = int(sheet_name)
        sheet = workbook[sheet_name]
        for excel_row in sheet.iter_rows(min_row=2, max_col=14, values_only=True):
            stat_date = parse_stat_date(excel_row[0], sheet_year)
            if stat_date is None:
                continue

            graphics = as_int(excel_row[3]) or 0
            plastics = as_int(excel_row[4]) or 0
            seats = as_int(excel_row[5]) or 0
            fitting = as_int(excel_row[6]) or 0
            without_fitting = as_int(excel_row[9])
            if without_fitting is None:
                without_fitting = graphics + plastics + seats

            after_weekend = as_int(excel_row[10])
            with_fitting = as_int(excel_row[11])
            if with_fitting is None:
                with_fitting = without_fitting + fitting

            day_code = str(excel_row[1]).strip().upper() if excel_row[1] else DAY_CODES[stat_date.isoweekday()]
            iso_week = as_int(excel_row[2]) or stat_date.isocalendar().week

            rows.append(
                {
                    "stat_date": stat_date.isoformat(),
                    "stat_year": sheet_year,
                    "iso_week": iso_week,
                    "day_code": day_code,
                    "graphics_count": graphics,
                    "plastics_count": plastics,
                    "seat_count": seats,
                    "fitting_count": fitting,
                    "products_without_fitting": without_fitting,
                    "after_weekend_count": after_weekend,
                    "products_with_fitting": with_fitting,
                }
            )

    return rows


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract daily Year-to-Year rows from WeeklyStat.xlsx.")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    rows = extract_rows(args.workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps({"source": str(args.workbook), "rows": rows}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(f"extracted_daily_rows={len(rows)}")


if __name__ == "__main__":
    main()
