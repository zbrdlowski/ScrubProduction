from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"E:\volume_0\darkscrub")
OUT = ROOT / "output" / "docs"
OUT.mkdir(parents=True, exist_ok=True)
DOCX_PATH = OUT / "multishipping_pracovny_postup.docx"
ASSETS = ROOT / "output" / "manual_assets"
ORDER_LIST_SCREENSHOT = ASSETS / "multishipping-order-list.png"
MULTISHIPPING_MODAL_SCREENSHOT = ASSETS / "multishipping-modal.png"
FEDEX_PREVIEW_SCREENSHOT = ASSETS / "multishipping-fedex-preview.png"

FONT = "Aptos"
TEAL = "167C80"
TEAL_LIGHT = "EAF6F6"
NAVY = "263746"
LIGHT = "F3F6F8"
BORDER = "D9D9D9"
MUTED = RGBColor(85, 95, 105)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=150, bottom=120, end=150):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), size)
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr_text, fld_char_2])


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(30, 36, 42)
    normal.paragraph_format.space_after = Pt(5.5)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = FONT
    title._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    title._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(7)

    for style_name, size, before, after in (
        ("Heading 1", 15, 14, 6),
        ("Heading 2", 11.5, 9, 3),
    ):
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_bullet(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.first_line_indent = Cm(-0.25)
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_run_font(r, bold=True)
        r = p.add_run(text[len(bold_lead):])
        set_run_font(r)
    else:
        set_run_font(p.add_run(text))
    return p


def add_numbered_step(doc, number, title, body):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.75)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(f"{number}.  {title}")
    set_run_font(r, size=10.7, bold=True, color=RGBColor(22, 124, 128))
    r = p.add_run(f"  {body}")
    set_run_font(r, size=10.5)
    return p


def add_figure(doc, image_path, caption, alt_text, width_cm=17.2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(image_path), width=Cm(width_cm))
    shape._inline.docPr.set("descr", alt_text)

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    cap.paragraph_format.keep_together = True
    r = cap.add_run(caption)
    set_run_font(r, size=8.8, color=MUTED)
    r.italic = True
    return p, cap


def add_label_table(doc):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(5.0)
    table.columns[1].width = Cm(12.0)
    headers = ["Označenie v tabuľke", "Význam"]
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        r = p.add_run(text)
        set_run_font(r, size=9.5, bold=True, color=RGBColor(255, 255, 255))
    rows = [
        ("Tyrkysová ikonka s počtom", "Otvorí výber objednávok pre nový multishipping."),
        ("MULTI · 3 orders", "Hlavná objednávka skupiny. Iba táto objednávka sa odošle do FedEx CSV."),
        ("MULTI → SO21490", "Pridružená objednávka. Odošle sa v rovnakej krabici ako uvedená hlavná objednávka."),
    ]
    for idx, (label, meaning) in enumerate(rows, start=1):
        cells = table.add_row().cells
        if idx % 2 == 0:
            for cell in cells:
                set_cell_shading(cell, LIGHT)
        for cell in cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        r = cells[0].paragraphs[0].add_run(label)
        set_run_font(r, size=9.5, bold=True, color=RGBColor(22, 124, 128))
        r = cells[1].paragraphs[0].add_run(meaning)
        set_run_font(r, size=9.5)
    set_repeat_table_header(table.rows[0])
    set_table_borders(table)
    return table


def add_state_table(doc):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(3.3), Cm(5.0), Cm(8.7)]
    for col, width in zip(table.columns, widths):
        col.width = width
    for i, text in enumerate(("Stav skupiny", "Čo znamená", "Čo môže pracovník urobiť")):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, TEAL)
        set_cell_margins(cell)
        r = cell.paragraphs[0].add_run(text)
        set_run_font(r, size=9.2, bold=True, color=RGBColor(255, 255, 255))
    rows = [
        ("Draft", "Skupina je uložená, CSV ešte nebolo vytvorené.", "Upraviť hlavný výber, pridať alebo odobrať objednávku, prípadne skupinu zrušiť."),
        ("Exported locked", "Skupina bola vložená do FedEx CSV a je zamknutá.", "Bežne už nič neupravovať. Pri chybe odomknúť skupinu, nepoužiť staré CSV a vytvoriť nové."),
        ("Shipped", "EOD import pridelil tracking a zásielku uzavrel.", "Skontrolovať tracking pri všetkých objednávkach. Skupina už nie je určená na úpravu."),
    ]
    for idx, row in enumerate(rows, start=1):
        cells = table.add_row().cells
        if idx % 2 == 1:
            for cell in cells:
                set_cell_shading(cell, TEAL_LIGHT)
        for col_idx, text in enumerate(row):
            cell = cells[col_idx]
            set_cell_margins(cell, top=130, bottom=130)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            r = cell.paragraphs[0].add_run(text)
            set_run_font(r, size=9.1, bold=(col_idx == 0))
    set_repeat_table_header(table.rows[0])
    set_table_borders(table)
    return table


doc = Document()
configure_styles(doc)

section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.7)
section.bottom_margin = Cm(1.6)
section.left_margin = Cm(1.8)
section.right_margin = Cm(1.8)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.add_run("Multishipping hotových objednávok")

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(11)
r = subtitle.add_run("Pracovný postup pre expedíciu")
set_run_font(r, size=12.5, bold=True, color=RGBColor(22, 124, 128))

intro = doc.add_paragraph()
intro.paragraph_format.space_after = Pt(9)
r = intro.add_run("Multishipping použite vtedy, keď viac hotových objednávok toho istého zákazníka balíte do jednej krabice s jedným tracking number. ")
set_run_font(r, size=10.8, bold=True)
r = intro.add_run("Systém odošle do FedExu iba jednu hlavnú objednávku a po importe EOD pridelí rovnaké sledovanie všetkým objednávkam v skupine.")
set_run_font(r, size=10.8)

doc.add_heading("Postup vytvorenia multishippingu", level=1)
add_numbered_step(doc, 1, "Otvorte Ready for Label", "Pracujte na stránke Orders so statusom Ready for Label.")
add_numbered_step(doc, 2, "Kliknite na tyrkysovú ikonku", "Ikonka s počtom pri mene zákazníka otvorí okno Multishipping.")
add_numbered_step(doc, 3, "Vyberte objednávky v krabici", "Štvorčekom vľavo zaškrtnite iba objednávky, ktoré skutočne vložíte do rovnakej krabice.")
add_numbered_step(doc, 4, "Určte hlavnú objednávku", "Kruhovým prepínačom Master napravo označte objednávku, ktorá sa použije vo FedEx CSV.")
add_numbered_step(doc, 5, "Skontrolujte adresy", "Všetky vybrané objednávky musia patriť tomu istému zákazníkovi a musia mať správnu spoločnú doručovaciu adresu.")
add_numbered_step(doc, 6, "Uložte skupinu", "Kliknite na Save multishipping. Objednávky zostanú v stave Ready for Label, ale systém ich označí ako jednu zásielku.")
add_numbered_step(doc, 7, "Skontrolujte označenia", "Hlavná objednávka dostane označenie MULTI s počtom objednávok; ostatné ukážu, s ktorou hlavnou objednávkou sa odosielajú.")

doc.add_page_break()
doc.add_heading("Vizuálna ukážka", level=1)
add_figure(
    doc,
    ORDER_LIST_SCREENSHOT,
    "Obrázok 1  Hlavná objednávka MULTI 4 a tri pridružené objednávky smerujúce k SO21490.",
    "Zoznam Ready for Label s hlavnou a pridruženými multishipping objednávkami.",
)
add_figure(
    doc,
    MULTISHIPPING_MODAL_SCREENSHOT,
    "Obrázok 2  Zaškrtnuté objednávky patria do jednej krabice; prepínač Master určuje riadok pre FedEx CSV.",
    "Multishipping modal so štyrmi objednávkami a zvolenou Master objednávkou.",
    width_cm=16.4,
)

doc.add_page_break()

doc.add_heading("Označenia v zozname objednávok", level=1)
add_label_table(doc)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(8)
r = p.add_run("Dôležité: ")
set_run_font(r, bold=True, color=RGBColor(22, 124, 128))
r = p.add_run("Pred vytvorením FedEx CSV musia byť všetky plánované multishipping skupiny uložené. Toto pravidlo platí na oboch pracoviskách.")
set_run_font(r, bold=True)

doc.add_heading("Export do FedExu", level=1)
add_bullet(doc, "Samostatná objednávka sa zobrazí v CSV náhľade normálne.")
add_bullet(doc, "Z multishipping skupiny sa v náhľade zobrazí iba hlavná objednávka.")
add_bullet(doc, "Pridružené objednávky sa do FedEx CSV nepridajú ako samostatné zásielky.")
add_bullet(doc, "Pri hlavnom riadku systém pripraví spoločné typy produktov, hodnotu, poistenie a odhad hmotnosti celej krabice.")
add_bullet(doc, "V náhľade skontrolujte a podľa skutočnosti upravte najmä hmotnosť, colnú hodnotu, poistenie, službu a adresu.")
add_bullet(doc, "FedEx referencia musí zostať ukončená číslom hlavnej objednávky, aby ju EOD import správne našiel.")
add_figure(
    doc,
    FEDEX_PREVIEW_SCREENSHOT,
    "Obrázok 3  FedEx CSV náhľad obsahuje iba hlavnú objednávku SO21490 označenú MULTI 4 orders.",
    "FedEx CSV náhľad s jedným hlavným multishipping riadkom a upraviteľnými hodnotami.",
)

doc.add_heading("Import EOD", level=1)
intro2 = doc.add_paragraph()
r = intro2.add_run("Po vytvorení zásielok vo FedEx Stratus nahrajte EOD report späť cez Import EOD. ")
set_run_font(r, bold=True)
r = intro2.add_run("Keď systém nájde hlavnú objednávku, automaticky spracuje celú skupinu.")
set_run_font(r)
for text in (
    "Rovnaké tracking number sa pridelí hlavnej aj všetkým pridruženým objednávkam.",
    "Všetky objednávky dostanú rovnaký dátum odoslania a status Shipped.",
    "Aktivita sa zapíše ku každej objednávke samostatne.",
    "Ak sa nepodarí spracovať jednu objednávku, systém vráti späť celú skupinu. Nevznikne čiastočne odoslaná zásielka.",
    "Opakovaný import rovnakého EOD nevytvorí duplicitné trackingy.",
    "Ak skupina obsahuje eBay objednávky, systém ich všetky pridá aj do eBay tracking exportu.",
):
    add_bullet(doc, text)

doc.add_heading("Stavy skupiny a opravy", level=1)
add_state_table(doc)

doc.add_heading("Najdôležitejšie pravidlá", level=1)
rules = [
    "Objednávkam ručne nemeňte status kvôli členstvu v multishippingu. Členstvo je samostatná informácia.",
    "Jedna objednávka môže patriť iba do jednej multishipping skupiny.",
    "Ak systém oznámi, že objednávka už bola exportovaná, nevytvárajte druhý štítok. Najprv skontrolujte export na druhom pracovisku.",
    "Po odomknutí skupiny nepoužívajte staré FedEx CSV. Vytvorte nový export.",
    "Po importe EOD skontrolujte tracking minimálne na hlavnej a jednej pridruženej objednávke.",
]
for rule in rules:
    add_bullet(doc, rule)

for sec in doc.sections:
    footer = sec.footer
    footer.distance = Cm(0.7)
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("DARKSCRUB  |  Multishipping pre expedíciu  |  Strana ")
    set_run_font(r, size=8.5, color=MUTED)
    add_page_number(p)

props = doc.core_properties
props.title = "Multishipping hotových objednávok"
props.subject = "Pracovný postup pre expedíciu"
props.author = "DARKSCRUB"
props.keywords = "multishipping, expedícia, FedEx, EOD"

doc.save(DOCX_PATH)
print(DOCX_PATH)
